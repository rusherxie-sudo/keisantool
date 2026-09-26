"""Generate editable, macro-free monthly workbooks. Requires openpyxl 3.1.x.
PYTHONPATH=/path/to/openpyxl-runtime python3 scripts/generate-workbooks.py
Formulas recalculate when opened in Excel / compatible spreadsheet software.
"""
from pathlib import Path
from datetime import date
import calendar
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties


def build(year, output):
    wb = Workbook()
    wb.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode='auto')
    guide = wb.active
    guide.title = '使い方'
    notes = [f'{year}年 勤怠管理テンプレート（無料・マクロなし）',
             '1〜12月の各シートへ出勤・退勤を時刻（例 09:00）で入力します。',
             '翌日退勤は1、当日退勤は0。休憩は整数の分で入力してください。',
             '黄色のセルを編集します。出退勤が両方空欄の日は集計しません。',
             '要修正があると月合計を表示しません。実働は分単位、最大24時間です。',
             '実働の単純合計です。残業・深夜割増・法定休憩・給与は計算しません。',
             '翌日勤務は開始日の月に全時間を計上します。',
             '数式はExcelなどで開くと再計算されます。印刷前に日付・結果を確認してください。',
             '無料で編集・社内配布できます。計算式のセルを変更した場合は結果をご確認ください。',
             '配布元 https://keisantool.com/kinmu-jikan/template/',
             'ブラウザ計算 https://keisantool.com/kinmu-jikan/']
    for n in notes: guide.append([n])
    guide.column_dimensions['A'].width = 115
    for row in guide: row[0].alignment = Alignment(wrap_text=True); guide.row_dimensions[row[0].row].height = 28
    for month in range(1,13):
        ws = wb.create_sheet(f'{month}月')
        ws.append([f'{year}年{month}月 勤怠表']); ws.merge_cells('A1:H1')
        ws.append(['黄色のセルに入力。翌日=1／当日=0、休憩は分。要修正があれば合計を表示しません。']);ws.merge_cells('A2:H2')
        ws.append(['日付','曜日','出勤','退勤','翌日(0/1)','休憩(分)','実働','確認'])
        end = calendar.monthrange(year,month)[1]+3
        for day in range(1,end-2):
            r=day+3; d=date(year,month,day)
            ws.append([d,'月火水木金土日'[d.weekday()],None,None,0,0])
            # Round time differences to integer minutes before validating/adding.
            span=f'ROUND((D{r}-C{r}+E{r})*1440,0)'
            invalid=f'OR(COUNT(C{r}:F{r})<>4,C{r}<0,C{r}>=1,D{r}<0,D{r}>=1,AND(E{r}<>0,E{r}<>1),F{r}<0,F{r}<>INT(F{r}),{span}<0,{span}>1440,F{r}>{span})'
            ws[f'H{r}']=f'=IF(AND(C{r}="",D{r}=""),"",IFERROR(IF({invalid},"要修正","OK"),"要修正"))'
            ws[f'G{r}']=f'=IF(H{r}="OK",({span}-F{r})/1440,"")'
            ws[f'A{r}'].number_format='m/d'
            for col in ['C','D']:ws[f'{col}{r}'].number_format='hh:mm'
            ws[f'G{r}'].number_format='[h]:mm'
            for col in ['C','D','E','F']:ws[f'{col}{r}'].fill=PatternFill('solid',fgColor='FFF2CC')
        total=end+2
        ws[f'A{total}']='月合計';ws[f'G{total}']=f'=IF(COUNTIF(H4:H{end},"要修正")>0,"要修正",SUM(G4:G{end}))';ws[f'G{total}'].number_format='[h]:mm'
        ws[f'A{total+1}']='小数時間';ws[f'G{total+1}']=f'=IF(COUNTIF(H4:H{end},"要修正")>0,"要修正",ROUND(SUM(G4:G{end})*24,2))'
        ws[f'A{total+2}']='勤務日数';ws[f'G{total+2}']=f'=IF(COUNTIF(H4:H{end},"要修正")>0,"要修正",COUNT(G4:G{end}))'
        ws[f'A{total+4}']='計算ツール https://keisantool.com/kinmu-jikan/'
        ws.merge_cells(start_row=total+4,start_column=1,end_row=total+4,end_column=8)
        for kind,low,high,area in [('time',0,'0.999999',f'C4:D{end}'),('whole',0,1,f'E4:E{end}'),('whole',0,1440,f'F4:F{end}')]:
            dv=DataValidation(type=kind,operator='between',formula1=low,formula2=high,allow_blank=True)
            dv.showErrorMessage=True;dv.errorTitle='入力を確認してください';dv.error='時刻、翌日0/1、休憩の整数分を確認してください。';ws.add_data_validation(dv);dv.add(area)
        for row in ws:
            for cell in row:cell.font=Font(name='Yu Gothic',size=10);cell.alignment=Alignment(vertical='center')
        for cell in ws[3]:cell.fill=PatternFill('solid',fgColor='DDEBF7');cell.font=Font(name='Yu Gothic',bold=True)
        for col in 'ABCDEFGH':ws.column_dimensions[col].width=15
        for r in range(1,total+5):ws.row_dimensions[r].height=20
        ws.freeze_panes='C4';ws.print_title_rows='1:3';ws.print_area=f'A1:H{total+4}'
        ws.page_setup.orientation='portrait';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToHeight=1;ws.page_setup.fitToWidth=1;ws.sheet_properties.pageSetUpPr.fitToPage=True
    wb.save(output)

if __name__=='__main__':
    out=Path('public/downloads');out.mkdir(exist_ok=True,parents=True)
    for year in [2026,2027]:build(year,out/f'kinmu-template-{year}.xlsx')
