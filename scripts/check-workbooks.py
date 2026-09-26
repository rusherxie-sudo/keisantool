"""Optional workbook formula check: requires openpyxl and formulas (external to app build)."""
from pathlib import Path
from tempfile import TemporaryDirectory
from datetime import time
import calendar
from openpyxl import load_workbook
import formulas

for year in [2026,2027]:
    w=load_workbook(f'public/downloads/kinmu-template-{year}.xlsx')
    assert len(w.sheetnames)==13
    for month in range(1,13):
        s=w[f'{month}月']
        dates=[r[0].value for r in s.iter_rows(min_row=4) if hasattr(r[0].value,'month')]
        assert len(dates)==calendar.monthrange(year,month)[1]
        assert all(d.day==i+1 and d.month==month and d.year==year for i,d in enumerate(dates))
        assert s.page_setup.fitToWidth==1 and s.page_setup.fitToHeight==1

with TemporaryDirectory() as tmp:
    w=load_workbook('public/downloads/kinmu-template-2026.xlsx')
    for name in list(w.sheetnames):
        if name!='9月': del w[name]
    s=w['9月']
    s['C4']=time(9);s['D4']=time(18);s['F4']=60
    s['C5']=time(22);s['D5']=time(6);s['E5']=1;s['F5']=45
    def calc():
        path=Path(tmp)/'fixture.xlsx';w.save(path)
        result=formulas.ExcelModel().loads(str(path)).finish().calculate()
        return {k.split('!')[-1]:v.value[0,0] for k,v in result.items() if hasattr(v,'value')}
    r=calc();assert abs(r['G4']*1440-480)<1e-8;assert abs(r['G5']*1440-435)<1e-8
    assert r['G36']==15.25 and r['G37']==2
    s['C6']=time(9);s['D6']=time(18);s['F6']=0
    r=calc();assert abs(r['G35']*24-24.25)<1e-8 # total must not wrap at 24h
    for start,end,nextday,pause in [(time(22),time(6),0,45),(time(9),None,0,0),(time(9),time(18),0,999),(time(9),time(18),0,0.5)]:
        s['C4']=start;s['D4']=end;s['E4']=nextday;s['F4']=pause
        r=calc();assert r['H4']=='要修正' and r['G35']=='要修正'
print('PASS: 24 monthly calendars, overnight calculations, >24h totals, partial times and invalid breaks; formula interpreter (not native Excel).')
